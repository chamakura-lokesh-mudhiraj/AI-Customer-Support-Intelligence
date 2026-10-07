from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from app.services.document_loader import (
    load_docx,
    load_document,
    load_pdf,
)


def test_load_docx(tmp_path: Path):
    document_path = tmp_path / "test_document.docx"

    document = Document()
    document.add_paragraph("Password reset instructions")
    document.add_paragraph(
        "Customers can reset their password from the login page."
    )
    document.save(document_path)

    text = load_docx(str(document_path))

    assert "Password reset instructions" in text
    assert "Customers can reset their password" in text


def test_load_pdf(tmp_path: Path):
    pdf_path = tmp_path / "test_document.pdf"

    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(100, 750, "Billing support instructions")
    pdf.drawString(100, 730, "Customers can request billing assistance.")
    pdf.save()

    text = load_pdf(str(pdf_path))

    assert "Billing support instructions" in text
    assert "Customers can request billing assistance." in text

def test_load_document_returns_metadata(tmp_path: Path):
    document_path = tmp_path / "support_guide.docx"

    document = Document()
    document.add_paragraph("Refund policy")
    document.add_paragraph(
        "Customers can request a refund within thirty days."
    )
    document.save(document_path)

    document_content = load_document(str(document_path))

    assert document_content.filename == "support_guide.docx"
    assert document_content.file_type == "docx"
    assert "Refund policy" in document_content.text
    assert document_content.character_count == len(document_content.text)

def test_load_txt(tmp_path):
    file_path = tmp_path / "refund_policy.txt"

    file_path.write_text(
        "Customers can request a refund within 30 days.",
        encoding="utf-8",
    )

    document = load_document(str(file_path))

    assert document.filename == "refund_policy.txt"
    assert document.file_type == "txt"
    assert document.text == (
        "Customers can request a refund within 30 days."
    )
    assert document.character_count == len(document.text)