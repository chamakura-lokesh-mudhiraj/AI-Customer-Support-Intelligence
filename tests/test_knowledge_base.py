from pathlib import Path

from docx import Document

from app.services.knowledge_base import ingest_document
from app.services.vector_store import VectorStore


def test_ingest_document(tmp_path: Path, monkeypatch):
    document_path = tmp_path / "refund_policy.docx"

    document = Document()
    document.add_paragraph("Refund Policy")
    document.add_paragraph(
        "Customers can request a refund within thirty days."
    )
    document.save(document_path)

    def fake_create_embedding(text: str) -> list[float]:
        if "refund" in text.lower():
            return [1.0, 0.0, 0.0]

        return [0.0, 1.0, 0.0]

    monkeypatch.setattr(
        "app.services.knowledge_base.create_embedding",
        fake_create_embedding,
    )

    monkeypatch.chdir(tmp_path)

    vector_store = VectorStore(
        collection_name="test_ingestion"
    )

    chunks = ingest_document(
        str(document_path),
        vector_store=vector_store,
    )

    assert len(chunks) > 0
    assert chunks[0].filename == "refund_policy.docx"
    assert "Refund Policy" in chunks[0].text

    results = vector_store.search(
        embedding=[1.0, 0.0, 0.0],
        n_results=1,
    )

    assert len(results["documents"]) == 1
    assert "refund" in results["documents"][0][0].lower()